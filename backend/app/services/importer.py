"""U9 上线迁移：主数据、期初库存、未结订单/采购和采购历史。

先“预检”（dry_run=True）返回逐行错误，不落库；确认无误再正式导入。整批在一个事务内，任何一行出错都不写入。
"""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.errors import BizError
from ..models import (Bom, BomLine, Location, Material, MaterialType, StockLot, TxnType, User, Warehouse, now,
                      Supplier, SalesOrder, PurchaseOrder, PurchaseOrderLine)
from .common import audit
from .inventory import stock_in
from . import orders as order_service

TEMPLATES = {
    "materials": ["编码", "名称", "规格", "单位", "类型", "图号", "通用件", "售后备件", "提前期天数",
                  "检验转运天数", "免BOM", "免图纸管控", "生产库下限", "生产库上限", "备件库下限", "备件库上限"],
    "boms": ["父件编码", "版本", "子件编码", "单位用量", "损耗率", "生效"],
    "stock": ["物料编码", "仓库编码", "库位编码", "批次号", "数量", "单价", "入库日期"],
    "suppliers": ["供应商编码", "供应商名称", "联系人", "电话", "是否外协厂"],
    "orders": ["订单号", "客户", "整机编码", "数量", "交期", "计划开工", "优先级", "迁移状态", "是否非标", "工位", "备注"],
    "open_pos": ["采购单号", "供应商编码", "物料编码", "未到数量", "单价", "需求日期", "承诺日期", "关联订单号", "采购员用户名"],
    "po_history": ["采购日期", "采购单号", "供应商编码", "物料编码", "数量", "单价"],
}
EXAMPLES = {
    "materials": ["WG-DJ-75", "三相异步电机", "YE3-280S-4 75kW", "台", "外购", "", "是", "是", 15, 1, "否", "是", 1, 6, 1, 3],
    "boms": ["CP-PYB900", "V1", "WG-DJ-75", 1, 0, "是"],
    "stock": ["WG-DJ-75", "WH01", "A-01", "期初-001", 3, 28500, "2026-09-01"],
    "suppliers": ["GYS-001", "示例供应商", "张三", "13800000000", "否"],
    "orders": ["OLD-SO-001", "示例客户", "CP-PYB900", 1, "2026-12-01", "2026-11-20", 3, "已下达", "否", "A工位", "旧U8迁移"],
    "open_pos": ["OLD-PO-001", "GYS-001", "WG-DJ-75", 2, 28500, "2026-11-20", "2026-11-18", "OLD-SO-001", "caigou"],
    "po_history": ["2026-08-01", "HIS-PO-001", "GYS-001", "WG-DJ-75", 3, 28000],
}
TYPE_MAP = {v: k for k, v in MaterialType.LABELS.items()} | {k: k for k in MaterialType.ALL}


def template(kind: str) -> bytes:
    if kind not in TEMPLATES:
        raise BizError("模板类型无效")
    wb = Workbook()
    ws = wb.active
    ws.title = kind
    ws.append(TEMPLATES[kind])
    ws.append(EXAMPLES[kind])
    for i, _ in enumerate(TEMPLATES[kind], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(content: bytes, kind: str) -> list[tuple[int, dict]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise BizError(f"无法读取 Excel：{e}") from e
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it, [])]
    need = TEMPLATES[kind][:2]
    if not all(h in head for h in need):
        raise BizError(f"表头不对，请使用模板（至少包含：{'、'.join(need)}）")
    out = []
    for n, row in enumerate(it, start=2):
        if row is None or all(v in (None, "") for v in row):
            continue
        out.append((n, {head[i]: (row[i].strip() if isinstance(row[i], str) else row[i])
                        for i in range(min(len(head), len(row))) if head[i]}))
    return out


def _dec(v, name: str, default: Decimal | None = None) -> Decimal:
    if v in (None, ""):
        if default is not None:
            return default
        raise ValueError(f"{name}不能为空")
    try:
        return Decimal(str(v))
    except InvalidOperation as e:
        raise ValueError(f"{name}不是数字：{v}") from e


def _yes(v) -> bool:
    return str(v).strip() in ("是", "Y", "y", "1", "true", "True", "TRUE")


def _date(v, name: str, required: bool = False) -> date | None:
    if v in (None, ""):
        if required:
            raise ValueError(f"{name}不能为空")
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError as e:
        raise ValueError(f"{name}日期格式错误：{v}") from e


def run(db: Session, user: User, kind: str, content: bytes, dry_run: bool = True) -> dict:
    rows = _rows(content, kind)
    errors: list[dict] = []
    stats = {"rows": len(rows), "created": 0, "updated": 0}
    try:
        if kind == "materials":
            _materials(db, rows, errors, stats)
        elif kind == "boms":
            _boms(db, rows, errors, stats)
        elif kind == "stock":
            _stock(db, user, rows, errors, stats)
        elif kind == "suppliers":
            _suppliers(db, rows, errors, stats)
        elif kind == "orders":
            _orders(db, user, rows, errors, stats)
        elif kind == "open_pos":
            _purchase_rows(db, rows, errors, stats, history=False)
        elif kind == "po_history":
            _purchase_rows(db, rows, errors, stats, history=True)
        else:
            raise BizError("导入类型无效")
        db.flush()
    except Exception:
        db.rollback()
        raise
    if errors or dry_run:
        db.rollback()
    else:
        audit(db, user.id, "IMPORT", kind, None, stats)
        db.flush()  # 不在这里提交：调用方把导入、审计和齐套重算放进同一个事务一次提交
    return {**stats, "errors": errors[:200], "error_count": len(errors), "committed": not errors and not dry_run}


def _materials(db, rows, errors, stats):
    seen = set()
    for n, r in rows:
        try:
            code = str(r.get("编码") or "").strip()
            if not code or not r.get("名称"):
                raise ValueError("编码和名称必填")
            if code in seen:
                raise ValueError("文件内编码重复")
            seen.add(code)
            t = TYPE_MAP.get(str(r.get("类型") or "").strip())
            if not t:
                raise ValueError("类型应为 整机/自制/外购/外协")
            m = db.scalar(select(Material).where(Material.code == code))
            new = m is None
            if new:
                m = Material(code=code, type=t)
                db.add(m)
            m.name, m.type = str(r["名称"]), t
            m.spec = str(r.get("规格") or "")
            m.unit = str(r.get("单位") or "件")
            if new or r.get("图号"):
                if not new and m.drawing_no and str(r.get("图号")) != m.drawing_no:
                    raise ValueError(f"图号与系统不一致（{m.drawing_no}），改图号请走技改换图号")
                m.drawing_no = str(r.get("图号") or "")
            m.is_common, m.is_spare = _yes(r.get("通用件")), _yes(r.get("售后备件"))
            m.lead_time_days = int(_dec(r.get("提前期天数"), "提前期", Decimal(7)))
            m.inbound_days = int(_dec(r.get("检验转运天数"), "检验转运天数", Decimal(1)))
            m.bom_exempt = _yes(r.get("免BOM"))
            m.drawing_exempt = _yes(r.get("免图纸管控"))
            m.safety_min = _dec(r.get("生产库下限"), "", Decimal(0))
            m.safety_max = _dec(r.get("生产库上限"), "", Decimal(0))
            m.spare_min = _dec(r.get("备件库下限"), "", Decimal(0))
            m.spare_max = _dec(r.get("备件库上限"), "", Decimal(0))
            stats["created" if new else "updated"] += 1
        except ValueError as e:
            errors.append({"row": n, "error": str(e)})


def _boms(db, rows, errors, stats):
    mats = {m.code: m for m in db.scalars(select(Material)).all()}
    groups: dict[tuple[str, str], list] = {}
    for n, r in rows:
        try:
            p, c = str(r.get("父件编码") or "").strip(), str(r.get("子件编码") or "").strip()
            if p not in mats:
                raise ValueError(f"父件 {p} 不存在，请先导入物料")
            if c not in mats:
                raise ValueError(f"子件 {c} 不存在，请先导入物料")
            if p == c:
                raise ValueError("父件和子件相同")
            groups.setdefault((p, str(r.get("版本") or "V1")), []).append(
                (n, mats[c].id, _dec(r.get("单位用量"), "单位用量"), _dec(r.get("损耗率"), "", Decimal(0)),
                 _yes(r.get("生效"))))
        except ValueError as e:
            errors.append({"row": n, "error": str(e)})
    for (p, ver), lines in groups.items():
        pm = mats[p]
        b = db.scalar(select(Bom).where(Bom.material_id == pm.id, Bom.version == ver))
        if b and b.status != "DRAFT":
            errors.append({"row": lines[0][0], "error": f"{p} {ver} 已生效/作废，不能覆盖，请换新版本号"})
            continue
        new = b is None
        if new:
            b = Bom(material_id=pm.id, version=ver, is_mature=True)
            db.add(b)
        b.lines.clear()
        ids = set()
        for i, (n, cid, q, s, _) in enumerate(lines, 1):
            if cid in ids:
                errors.append({"row": n, "error": "同一 BOM 子件重复"})
                continue
            ids.add(cid)
            b.lines.append(BomLine(seq=i * 10, child_id=cid, qty_per=q, scrap_rate=s))
        if any(x[4] for x in lines):
            for o in db.scalars(select(Bom).where(Bom.material_id == pm.id, Bom.status == "ACTIVE")).all():
                o.status = "OBSOLETE"
            b.status = "ACTIVE"
        stats["created" if new else "updated"] += 1


def _stock(db, user, rows, errors, stats):
    mats = {m.code: m for m in db.scalars(select(Material)).all()}
    whs = {w.code: w for w in db.scalars(select(Warehouse)).all()}
    for n, r in rows:
        try:
            m, w = mats.get(str(r.get("物料编码") or "").strip()), whs.get(str(r.get("仓库编码") or "").strip())
            if not m:
                raise ValueError("物料不存在")
            if not w:
                raise ValueError("仓库编码不存在")
            loc = None
            if r.get("库位编码"):
                loc = db.scalar(select(Location).where(Location.warehouse_id == w.id, Location.code == str(r["库位编码"])))
                if not loc:
                    raise ValueError("库位不存在")
            q = _dec(r.get("数量"), "数量")
            if q <= 0:
                raise ValueError("数量必须大于 0")
            lot = stock_in(db, material_id=m.id, warehouse_id=w.id, qty=q, txn_type=TxnType.ADJUST,
                           lot_no=str(r.get("批次号") or f"INIT-{n}"), location_id=loc.id if loc else None,
                           unit_cost=_dec(r.get("单价"), "", Decimal(0)), ref_type="INIT", user_id=user.id,
                           remark="期初导入")
            d = r.get("入库日期")
            if d:
                dt_ = d if isinstance(d, datetime) else datetime.fromisoformat(str(d)[:10])
                lot.received_at = lot.last_move_at = dt_
            stats["created"] += 1
        except ValueError as e:
            errors.append({"row": n, "error": str(e)})


def _suppliers(db, rows, errors, stats):
    seen = set()
    for n, r in rows:
        try:
            code, name = str(r.get("供应商编码") or "").strip(), str(r.get("供应商名称") or "").strip()
            if not code or not name:
                raise ValueError("供应商编码和名称必填")
            if code in seen:
                raise ValueError("文件内供应商编码重复")
            seen.add(code)
            s = db.scalar(select(Supplier).where(Supplier.code == code))
            new = s is None
            if new:
                s = Supplier(code=code, name=name)
                db.add(s)
            s.name, s.contact, s.phone = name, str(r.get("联系人") or ""), str(r.get("电话") or "")
            s.is_outsourcer, s.is_active = _yes(r.get("是否外协厂")), True
            stats["created" if new else "updated"] += 1
        except ValueError as e:
            errors.append({"row": n, "error": str(e)})


def _orders(db, user, rows, errors, stats):
    mats = {m.code: m for m in db.scalars(select(Material).where(Material.type == MaterialType.PRODUCT)).all()}
    seen = set()
    for n, r in rows:
        try:
            no = str(r.get("订单号") or "").strip()
            if not no or no in seen:
                raise ValueError("订单号必填且文件内不能重复")
            seen.add(no)
            if db.scalar(select(SalesOrder).where(SalesOrder.order_no == no)):
                raise ValueError("订单号已存在")
            m = mats.get(str(r.get("整机编码") or "").strip())
            if not m:
                raise ValueError("整机编码不存在或不是整机")
            state = str(r.get("迁移状态") or "已录入").strip()
            if state not in ("已录入", "已下达"):
                raise ValueError("迁移状态仅支持 已录入/已下达；装配中订单须先核对在制领料后专项迁移")
            o = order_service.create_order(db, user, {"order_no": no, "customer": str(r.get("客户") or "").strip(),
                "product_id": m.id, "qty": _dec(r.get("数量"), "数量"), "due_date": _date(r.get("交期"), "交期", True),
                "planned_start": _date(r.get("计划开工"), "计划开工"), "priority": int(_dec(r.get("优先级"), "优先级", Decimal(3))),
                "is_nonstandard": _yes(r.get("是否非标")), "station": str(r.get("工位") or ""),
                "remark": str(r.get("备注") or "")})
            if state == "已下达":
                order_service.release(db, user, o.id)
            stats["created"] += 1
        except (ValueError, BizError) as e:
            errors.append({"row": n, "error": e.message if isinstance(e, BizError) else str(e)})


def _purchase_rows(db, rows, errors, stats, history: bool):
    suppliers = {x.code: x for x in db.scalars(select(Supplier)).all()}
    mats = {x.code: x for x in db.scalars(select(Material)).all()}
    orders = {x.order_no: x for x in db.scalars(select(SalesOrder)).all()}
    users = {x.username: x for x in db.scalars(select(User)).all()}
    grouped = {}
    first_row = {}
    for n, r in rows:
        no = str(r.get("采购单号") or "").strip()
        if not no:
            errors.append({"row": n, "error": "采购单号必填"})
            continue
        grouped.setdefault(no, []).append((n, r)); first_row.setdefault(no, n)
    for no, items in grouped.items():
        try:
            if db.scalar(select(PurchaseOrder).where(PurchaseOrder.po_no == no)):
                raise ValueError("采购单号已存在")
            supplier_codes = {str(r.get("供应商编码") or "").strip() for _, r in items}
            if len(supplier_codes) != 1 or next(iter(supplier_codes)) not in suppliers:
                raise ValueError("同一采购单必须对应一个已存在供应商")
            supplier = suppliers[next(iter(supplier_codes))]
            buyer = None if history else users.get(str(items[0][1].get("采购员用户名") or "").strip())
            po = PurchaseOrder(po_no=no, supplier_id=supplier.id, buyer_id=buyer.id if buyer else None,
                               status="CLOSED" if history else "APPROVED",
                               order_date=_date(items[0][1].get("采购日期"), "采购日期") or date.today(),
                               remark="旧U8采购历史迁移" if history else "旧U8未结采购迁移")
            db.add(po); db.flush()
            seen_mids = set()
            for i, (n, r) in enumerate(items, 1):
                m = mats.get(str(r.get("物料编码") or "").strip())
                if not m:
                    raise ValueError(f"第 {n} 行物料不存在")
                if m.id in seen_mids:
                    raise ValueError(f"第 {n} 行同一采购单物料重复，请先合并")
                seen_mids.add(m.id)
                q = _dec(r.get("数量") if history else r.get("未到数量"), "数量")
                if q <= 0:
                    raise ValueError(f"第 {n} 行数量必须大于0")
                so = orders.get(str(r.get("关联订单号") or "").strip()) if not history else None
                po.lines.append(PurchaseOrderLine(line_no=i * 10, material_id=m.id, drawing_no=m.drawing_no,
                    qty=q, price=_dec(r.get("单价"), "单价", Decimal(0)),
                    need_date=None if history else _date(r.get("需求日期"), "需求日期"),
                    promised_date=None if history else _date(r.get("承诺日期"), "承诺日期"),
                    arrived_qty=q if history else Decimal(0), received_qty=q if history else Decimal(0),
                    status="CLOSED" if history else "OPEN", order_id=so.id if so else None))
            stats["created"] += len(items)
        except ValueError as e:
            errors.append({"row": first_row[no], "error": str(e)})


__all__ = ["template", "run", "TEMPLATES", "StockLot", "now"]
